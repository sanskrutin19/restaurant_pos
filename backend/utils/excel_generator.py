from openpyxl import Workbook
from openpyxl.styles import Font, Fill, PatternFill, Border, Side, Alignment, Color
from openpyxl.chart import PieChart, BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from datetime import datetime
import io

class ExcelGenerator:
    BLUE = "2563EB"
    LIGHT_BLUE = "DBEAFE"
    GREEN = "059669"
    LIGHT_GREEN = "D1FAE5"
    ORANGE = "EA580C"
    RED = "DC2626"
    GRAY = "F9FAFB"
    WHITE = "FFFFFF"

    @staticmethod
    def generate_sales_excel(data: dict, start_date: datetime, end_date: datetime, restaurant_name: str):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales Report"

        # Styles
        header_font = Font(name='Arial', size=16, bold=True, color=ExcelGenerator.WHITE)
        header_fill = PatternFill(start_color=ExcelGenerator.BLUE, end_color=ExcelGenerator.BLUE, fill_type="solid")
        section_font = Font(name='Arial', size=14, bold=True, color=ExcelGenerator.WHITE)
        table_header_fill = PatternFill(start_color=ExcelGenerator.LIGHT_BLUE, end_color=ExcelGenerator.LIGHT_BLUE, fill_type="solid")
        border = Border(left=Side(style='thin', color="E5E7EB"), 
                        right=Side(style='thin', color="E5E7EB"), 
                        top=Side(style='thin', color="E5E7EB"), 
                        bottom=Side(style='thin', color="E5E7EB"))
        center_align = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')
        right_align = Alignment(horizontal='right', vertical='center')

        # 1. HEADER SECTION (Rows 1-4)
        ws.merge_cells('A1:F1')
        ws['A1'] = restaurant_name.upper()
        ws['A1'].font = header_font
        ws['A1'].fill = header_fill
        ws['A1'].alignment = center_align

        ws.merge_cells('A2:F2')
        ws['A2'] = "SALES REPORT"
        ws['A2'].font = Font(size=14, bold=True, color=ExcelGenerator.WHITE)
        ws['A2'].fill = header_fill
        ws['A2'].alignment = center_align

        ws.merge_cells('A3:F3')
        ws['A3'] = f"Period: {start_date.strftime('%d/%m/%Y')} to {end_date.strftime('%d/%m/%Y')}"
        ws['A3'].font = Font(size=12, color=ExcelGenerator.WHITE)
        ws['A3'].fill = header_fill
        ws['A3'].alignment = center_align

        ws.merge_cells('A4:F4')
        ws['A4'] = f"Generated On: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws['A4'].font = Font(size=10, italic=True, color=ExcelGenerator.WHITE)
        ws['A4'].fill = header_fill
        ws['A4'].alignment = center_align

        # 2. SUMMARY SECTION (Starting Row 6)
        curr_row = 6
        ws.merge_cells(f'A{curr_row}:B{curr_row}')
        ws[f'A{curr_row}'] = "EXECUTIVE SUMMARY"
        ws[f'A{curr_row}'].font = section_font
        ws[f'A{curr_row}'].fill = header_fill
        
        curr_row += 1
        headers = ["Metric", "Value"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=curr_row, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = table_header_fill
            cell.border = border

        s = data['summary']
        summary_rows = [
            ("Total Revenue", s['total_revenue'], ExcelGenerator.GREEN),
            ("Total Orders", s['total_orders'], ExcelGenerator.BLUE),
            ("CGST", s['cgst'], ExcelGenerator.ORANGE),
            ("SGST", s['sgst'], ExcelGenerator.ORANGE),
            ("Total Tax", s['total_tax'], ExcelGenerator.RED),
            ("Total Discount", s['discount'], ExcelGenerator.RED)
        ]

        for metric, value, color in summary_rows:
            curr_row += 1
            ws.cell(row=curr_row, column=1, value=metric).border = border
            val_cell = ws.cell(row=curr_row, column=2, value=value)
            val_cell.border = border
            val_cell.font = Font(color=color, bold=True)
            if isinstance(value, (int, float)) and "Orders" not in metric:
                val_cell.number_format = '"₹"#,##0.00'

        curr_row += 2

        # 3. ORDERS SECTION
        ws.merge_cells(f'A{curr_row}:G{curr_row}')
        ws[f'A{curr_row}'] = "ORDERS HISTORY"
        ws[f'A{curr_row}'].font = section_font
        ws[f'A{curr_row}'].fill = header_fill
        
        curr_row += 1
        headers = ["Date", "Order ID", "Customer", "Items", "Amount", "Payment", "Status"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=curr_row, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = table_header_fill
            cell.border = border

        for i, o in enumerate(data['orders']):
            curr_row += 1
            row_data = [o['date'], o['id'], o['customer'], o['items'], o['amount'], o['payment_mode'], o['status']]
            fill = PatternFill(start_color=ExcelGenerator.GRAY, end_color=ExcelGenerator.GRAY, fill_type="solid") if i % 2 == 1 else None
            
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=curr_row, column=col, value=val)
                cell.border = border
                if fill: cell.fill = fill
                if col == 5: # Amount
                    cell.font = Font(color=ExcelGenerator.GREEN, bold=True)
                    cell.number_format = '"₹"#,##0.00'
                if col == 7: # Status
                    cell.font = Font(bold=True, color="059669")
                    cell.fill = PatternFill(start_color=ExcelGenerator.LIGHT_GREEN, end_color=ExcelGenerator.LIGHT_GREEN, fill_type="solid")

        curr_row += 2

        curr_row += 2
        
        # 4. ITEM-WISE SALES
        ws.merge_cells(f'A{curr_row}:C{curr_row}')
        ws[f'A{curr_row}'] = "ITEM-WISE SALES"
        ws[f'A{curr_row}'].font = section_font
        ws[f'A{curr_row}'].fill = header_fill
        
        curr_row += 1
        headers = ["Item Name", "Qty Sold", "Total Revenue"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=curr_row, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = table_header_fill
            cell.border = border

        item_start_row = curr_row + 1
        for i, item in enumerate(data['item_wise']):
            curr_row += 1
            row_data = [item['item_name'], item['quantity_sold'], item['total_revenue']]
            fill = PatternFill(start_color=ExcelGenerator.GRAY, end_color=ExcelGenerator.GRAY, fill_type="solid") if i % 2 == 1 else None
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=curr_row, column=col, value=val)
                cell.border = border
                if fill: cell.fill = fill
                if col == 2: cell.alignment = center_align
                if col == 3: 
                    cell.font = Font(color=ExcelGenerator.GREEN)
                    cell.number_format = '"₹"#,##0.00'
        item_end_row = curr_row

        curr_row += 2

        curr_row += 2
        
        # 5. CATEGORY-WISE SALES
        ws.merge_cells(f'A{curr_row}:C{curr_row}')
        ws[f'A{curr_row}'] = "CATEGORY-WISE SALES"
        ws[f'A{curr_row}'].font = section_font
        ws[f'A{curr_row}'].fill = header_fill
        
        curr_row += 1
        headers = ["Category", "Qty Sold", "Total Revenue"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=curr_row, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = table_header_fill
            cell.border = border

        cat_start_row = curr_row + 1
        for i, cat in enumerate(data['category_wise']):
            curr_row += 1
            row_data = [cat['category'], cat['quantity_sold'], cat['total_revenue']]
            fill = PatternFill(start_color=ExcelGenerator.GRAY, end_color=ExcelGenerator.GRAY, fill_type="solid") if i % 2 == 1 else None
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=curr_row, column=col, value=val)
                cell.border = border
                if fill: cell.fill = fill
                if col == 3: 
                    cell.font = Font(color=ExcelGenerator.GREEN)
                    cell.number_format = '"₹"#,##0.00'
        cat_end_row = curr_row

        curr_row += 2
        
        # 6. PAYMENT BREAKDOWN
        ws.merge_cells(f'A{curr_row}:C{curr_row}')
        ws[f'A{curr_row}'] = "PAYMENT BREAKDOWN"
        ws[f'A{curr_row}'].font = section_font
        ws[f'A{curr_row}'].fill = header_fill
        
        curr_row += 1
        headers = ["Payment Mode", "Count", "Total Amount"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=curr_row, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = table_header_fill
            cell.border = border

        pay_start_row = curr_row + 1
        for i, pay in enumerate(data['payment_breakdown']):
            curr_row += 1
            mode = pay['mode']
            row_data = [mode, pay['count'], pay['total_amount']]
            
            # Specific mode coloring
            bg_color = ExcelGenerator.WHITE
            if "Cash" in mode: bg_color = ExcelGenerator.LIGHT_GREEN
            elif "UPI" in mode: bg_color = ExcelGenerator.LIGHT_BLUE
            elif "Card" in mode: bg_color = "FEF3C7" # Light Yellow

            fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=curr_row, column=col, value=val)
                cell.border = border
                cell.fill = fill
                if col == 3: cell.number_format = '"₹"#,##0.00'
        pay_end_row = curr_row

        curr_row += 2

        # 7. REVENUE TREND (For Line Chart)
        ws.merge_cells(f'A{curr_row}:B{curr_row}')
        ws[f'A{curr_row}'] = "DAILY REVENUE TREND"
        ws[f'A{curr_row}'].font = section_font
        ws[f'A{curr_row}'].fill = header_fill
        
        curr_row += 1
        headers = ["Date", "Revenue"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=curr_row, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = table_header_fill
            cell.border = border
            
        trend_start_row = curr_row + 1
        for i, row in enumerate(data['revenue_trend']):
            curr_row += 1
            ws.cell(row=curr_row, column=1, value=row['day']).border = border
            val_cell = ws.cell(row=curr_row, column=2, value=row['revenue'])
            val_cell.border = border
            val_cell.number_format = '"₹"#,##0.00'
        trend_end_row = curr_row

        # --- VISUALIZATION DASHBOARD (Column J) ---
        # 1. Line Chart: Revenue Trend
        if data['revenue_trend'] and trend_end_row >= trend_start_row:
            line = LineChart()
            line.title = "Daily Revenue Trend"
            line.style = 13
            line.y_axis.title = 'Revenue (₹)'
            line.x_axis.title = 'Date'
            line.width = 15
            line.height = 8
            
            data_ref = Reference(ws, min_col=2, min_row=trend_start_row-1, max_row=trend_end_row)
            cats_ref = Reference(ws, min_col=1, min_row=trend_start_row, max_row=trend_end_row)
            line.add_data(data_ref, titles_from_data=True)
            line.set_categories(cats_ref)
            ws.add_chart(line, "J2")

        # 2. Bar Chart: Categories
        if data['category_wise'] and cat_end_row >= cat_start_row:
            bar = BarChart()
            bar.title = "Revenue by Category"
            bar.y_axis.title = 'Revenue (₹)'
            bar.width = 15
            bar.height = 8
            
            data_ref = Reference(ws, min_col=3, min_row=cat_start_row-1, max_row=cat_end_row)
            cats_ref = Reference(ws, min_col=1, min_row=cat_start_row, max_row=cat_end_row)
            bar.add_data(data_ref, titles_from_data=True)
            bar.set_categories(cats_ref)
            ws.add_chart(bar, "J18")

        # 3. Pie Chart: Payments
        if data['payment_breakdown'] and pay_end_row >= pay_start_row:
            pie = PieChart()
            pie.title = "Payment Distribution"
            pie.width = 15
            pie.height = 8
            labels = Reference(ws, min_col=1, min_row=pay_start_row, max_row=pay_end_row)
            data_ref = Reference(ws, min_col=3, min_row=pay_start_row-1, max_row=pay_end_row)
            pie.add_data(data_ref, titles_from_data=True)
            pie.set_categories(labels)
            ws.add_chart(pie, "J34")

        # Auto-fit columns
        from openpyxl.utils import get_column_letter
        for col in ws.columns:
            max_length = 0
            column_id = col[0].column
            column_letter = get_column_letter(column_id)
            if column_id > 8: continue # Don't auto-fit chart columns
            for cell in col:
                try:
                    if cell.value:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                except: pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column_letter].width = max(adjusted_width, 15)

        # Freeze Panes (Header section and Metric/Mode names)
        ws.freeze_panes = 'B5'

        # Conditional Formatting for Revenue (Top items)
        # Using a rule to highlight high revenue in green
        green_fill = PatternFill(start_color=ExcelGenerator.LIGHT_GREEN, end_color=ExcelGenerator.LIGHT_GREEN, fill_type="solid")
        ws.conditional_formatting.add(f'C{item_start_row}:C{item_end_row}',
                                      CellIsRule(operator='greaterThan', formula=['1000'], font=Font(bold=True, color="000000"), fill=green_fill))

        # Save to buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
